#!/usr/bin/env python3
"""Hand-labelled column roles: the one measurement the rest of the suite lacks.

    python labels.py make                  # write reports/labels/column_roles.xlsx
    python labels.py make --pages 40       # choose how many pages go in the sheet
    python labels.py score                 # score the current code against it
    python labels.py score --json out.json

No other metric here can see a *wrong role*. `11.pdf` page 2 reported all four
stock roles transposed and `columns_mapped` read 4 of 4 before and after the
fix; every count in the benchmark was blind to it. A person marking which
column is opening, receipt, issue and closing - a minute or two a page, not a
transcription of every cell - is what makes that visible.

Labels are keyed on **where a column sits on the page** (its horizontal span,
within the vertical span of its table), never on column index. The same table
resolved 7 columns on one page and 8 on the next, so an index-keyed label
would silently point at a different column after a code change.

The sheet is pre-filled with the system's own guesses so the labeller corrects
rather than types. That makes accepting a wrong guess easy, which is why only
pages explicitly marked reviewed are ever scored.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent
CORPUS = ROOT / "JUNE"
RESULTS = ROOT / "reports" / "_results"
LABEL_DIR = ROOT / "reports" / "labels"
WORKBOOK = LABEL_DIR / "column_roles.xlsx"

#: Labels beyond the canonical roles, for columns that are not one clean column.
MERGED = "merged"              # one band holding two or more printed columns
NOT_A_COLUMN = "not_a_column"  # a band that is not a real column at all
EXTRA_LABELS = {
    MERGED: "This band holds two or more printed columns squeezed together "
            "(e.g. 'Pur SP', or cells like '109. 100.').",
    NOT_A_COLUMN: "Not a real column at all: a stray band of text, a page "
                  "number, a fragment of the address block.",
}

ROLE_MEANING = {
    "item_description": "Product / item name",
    "code": "Item or product code",
    "pack": "Pack size (10'S, 1X10, 100ML)",
    "unit": "Unit of measure",
    "rate": "Price per unit",
    "opening_qty": "Opening stock, quantity",
    "receipt_qty": "Purchases / received during the period, quantity",
    "issue_qty": "Sales / issued during the period, quantity",
    "closing_qty": "Closing stock / balance, quantity",
    "opening_value": "Opening stock, rupee value",
    "receipt_value": "Purchases / received, rupee value",
    "issue_value": "Sales / issued, rupee value",
    "closing_value": "Closing stock / balance, rupee value",
    "dump_qty": "Damaged / breakage / dump, quantity",
    "unknown": "A real column with no role in this list: batch, expiry, MRP, "
               "free qty, sales return, near-expiry, order qty, discount, etc.",
}

SHEET_COLUMNS = ["file", "page", "table", "column", "heading as printed",
                 "sample values", "system role", "correct role",
                 "x0", "x1", "y0", "y1"]


# ---------------------------------------------------------------------------
# Geometry
# ---------------------------------------------------------------------------


def _section_rows_span(section) -> tuple[float, float] | None:
    ys = [r.bbox for r in section.data_rows if r.bbox]
    if not ys:
        return None
    return (min(b[1] for b in ys), max(b[3] for b in ys))


def _column_span(section, column) -> tuple[float, float] | None:
    """Median horizontal span of a column's cells; its heading if it has none.

    An empty column still matters - `GENEX.pdf` prints a receipt-quantity column
    that is blank on the whole page - so the heading's box stands in for it.
    """
    left, right = [], []
    for row in section.data_rows:
        cell = row.cell_at(column.index)
        if cell is None or cell.is_blank or not cell.provenance.bbox:
            continue
        left.append(cell.provenance.bbox[0])
        right.append(cell.provenance.bbox[2])
    if not left:
        for row in section.header_rows:
            cell = row.cell_at(column.index)
            if cell is not None and not cell.is_blank and cell.provenance.bbox:
                left.append(cell.provenance.bbox[0])
                right.append(cell.provenance.bbox[2])
    if not left:
        return None
    left.sort()
    right.sort()
    return (left[len(left) // 2], right[len(right) // 2])


@dataclass
class PredictedColumn:
    table: int
    column: int
    heading: str
    samples: str
    role: str
    x: tuple[float, float]
    y: tuple[float, float]


def predicted_columns(document) -> list[PredictedColumn]:
    from schema import SectionKind

    out: list[PredictedColumn] = []
    table = 0
    for section in document.sections:
        if section.kind is not SectionKind.TABULAR:
            continue
        y = _section_rows_span(section)
        if y is None:
            continue
        table += 1
        for position, column in enumerate(sorted(section.columns, key=lambda c: c.index), 1):
            x = _column_span(section, column)
            if x is None:
                continue
            samples = []
            for row in section.data_rows:
                cell = row.cell_at(column.index)
                if cell is not None and not cell.is_blank:
                    samples.append(cell.raw_text.strip())
                if len(samples) == 4:
                    break
            out.append(PredictedColumn(
                table=table, column=position, heading=column.header_text,
                samples=" | ".join(samples), role=column.role.value, x=x, y=y))
    return out


def _overlap(a: tuple[float, float], b: tuple[float, float]) -> float:
    return min(a[1], b[1]) - max(a[0], b[0])


def match(label_x, label_y, candidates: list[PredictedColumn]) -> PredictedColumn | None:
    """The predicted column occupying the labelled position, if any does.

    Scored by overlap *relative to the combined span*, not absolute overlap. On
    `IMG-20260630-WA0015 .jpg` a narrow column at x 771-785 sits entirely inside
    its neighbour's span of 687-785, so both overlap it by exactly 14 and the
    wider neighbour won the tie - a label reported as a wrong role that the
    system had in fact got right.
    """
    best, best_score = None, 0.0
    for candidate in candidates:
        if _overlap(label_y, candidate.y) <= 0:
            continue
        overlap = _overlap(label_x, candidate.x)
        if overlap <= 0:
            continue
        union = max(label_x[1], candidate.x[1]) - min(label_x[0], candidate.x[0])
        score = overlap / union if union > 0 else 0.0
        if score > best_score:
            best, best_score = candidate, score
    return best


# ---------------------------------------------------------------------------
# make
# ---------------------------------------------------------------------------


def choose_pages(count: int) -> list[tuple[str, int, str]]:
    """Pick pages for breadth of layout first, depth second.

    Half image pages, half text-layer pages. Within each, one page per file
    before any file gets a second - a first attempt that ranked continuation
    pages first filled the image quota from three multi-page scans and left out
    every WhatsApp photo. A file's continuation page is still preferred over its
    first page, because that is where headings go missing and roles get
    inherited, which is where this project keeps finding wrong roles.
    """
    pages = []
    for path in sorted(RESULTS.glob("batch-*.json")):
        for page in json.loads(path.read_text(encoding="utf-8"))["pages"]:
            if page["error"] or not page["tabular_sections"] or not page["columns"]:
                continue
            pages.append(page)
    if not pages:
        raise SystemExit("no batch results found; run batch_test.py first")

    # What each page physically is. Walking files alphabetically let names that
    # start with digits use up the image quota before `IMG-...-WA` and
    # `WhatsApp Image ...` got a turn, so the image half is shared out across
    # these groups instead of across the alphabet.
    def group(page: dict) -> str:
        if page["source"] == "pdf_text":
            return "text-layer PDF"
        if page["kind"] == "pdf":
            return "scanned PDF"
        return "WhatsApp photo" if page["kind"] == "whatsapp" else "scanned photo"

    groups: dict[str, dict[str, list[dict]]] = defaultdict(lambda: defaultdict(list))
    for page in pages:
        groups[group(page)][page["file"]].append(page)
    for files in groups.values():
        for file_pages in files.values():
            # Continuation pages first, then page 1.
            file_pages.sort(key=lambda p: (p["page"] == 1, p["page"]))

    image_groups = [g for g in ("scanned PDF", "scanned photo", "WhatsApp photo") if g in groups]
    quotas: dict[str, int] = {"text-layer PDF": count - count // 2}
    for i, name in enumerate(image_groups):
        quotas[name] = count // 2 // len(image_groups) + (i < (count // 2) % len(image_groups))

    chosen: list[tuple[str, int, str]] = []
    for name, wanted in quotas.items():
        files = groups.get(name, {})
        taken = depth = 0
        while taken < wanted and any(len(v) > depth for v in files.values()):
            for file_name in sorted(files, key=str.lower):
                if taken >= wanted:
                    break
                if len(files[file_name]) > depth:
                    page = files[file_name][depth]
                    chosen.append((file_name, page["page"], page["source"]))
                    taken += 1
            depth += 1
    return sorted(chosen, key=lambda c: (c[0].lower(), c[1]))


def make(count: int) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from pipeline import process_file
    from schema import Role

    chosen = choose_pages(count)
    by_file: dict[str, list[int]] = defaultdict(list)
    for name, page, _ in chosen:
        by_file[name].append(page)

    book = Workbook()
    guide = book.active
    guide.title = "Instructions"
    for line in (
        "Column role labels",
        "",
        "One row per column the system found, on the pages listed in the Pages sheet.",
        "Open each source file from the JUNE folder beside this sheet.",
        "",
        "1. Read 'heading as printed' and 'sample values', and find the column on the page.",
        "2. 'system role' is what the system guessed. 'correct role' starts as a copy of it.",
        "3. If the guess is wrong, change 'correct role' using its dropdown (meanings in the Roles sheet).",
        "4. When every row of a page is right, put Y in that page's 'reviewed' cell in the Pages sheet.",
        "",
        "Only pages marked Y are scored. Do not mark a page Y without checking every row:",
        "the correct-role column is pre-filled, so an unchecked page would score the system against itself.",
        "",
        "Use 'unknown' for a real column with no role in the list (free qty, returns, near-expiry...).",
        f"Use '{MERGED}' when one row holds values from two printed columns.",
        f"Use '{NOT_A_COLUMN}' for a band that is not a column at all.",
        "",
        "Do not edit the file, page, table, column or the hidden x0/x1/y0/y1 columns.",
    ):
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=14)
    guide.column_dimensions["A"].width = 110

    roles_sheet = book.create_sheet("Roles")
    roles_sheet.append(["role", "meaning"])
    role_values = [r.value for r in Role] + [MERGED, NOT_A_COLUMN]
    for value in role_values:
        roles_sheet.append([value, ROLE_MEANING.get(value) or EXTRA_LABELS.get(value, "")])
    roles_sheet.column_dimensions["A"].width = 18
    roles_sheet.column_dimensions["B"].width = 90
    for cell in roles_sheet[1]:
        cell.font = Font(bold=True)

    pages_sheet = book.create_sheet("Pages")
    pages_sheet.append(["file", "page", "source", "columns", "reviewed"])
    for cell in pages_sheet[1]:
        cell.font = Font(bold=True)

    sheet = book.create_sheet("Labels", 1)
    sheet.append(SHEET_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    guess_fill = PatternFill("solid", fgColor="EDEDED")
    label_fill = PatternFill("solid", fgColor="FFF2CC")
    rows_written = 0
    for name in sorted(by_file, key=str.lower):
        print(f"  reading {name} ...", flush=True)
        result = process_file(CORPUS / name)
        for page_number in sorted(by_file[name]):
            page = result.pages[page_number - 1]
            if page.document is None:
                continue
            columns = predicted_columns(page.document)
            kind = next(k for n, p, k in chosen if n == name and p == page_number)
            pages_sheet.append([name, page_number, kind, len(columns), ""])
            for col in columns:
                sheet.append([name, page_number, col.table, col.column, col.heading,
                              col.samples, col.role, col.role,
                              round(col.x[0], 2), round(col.x[1], 2),
                              round(col.y[0], 2), round(col.y[1], 2)])
                rows_written += 1
                sheet.cell(row=sheet.max_row, column=7).fill = guess_fill
                sheet.cell(row=sheet.max_row, column=8).fill = label_fill

    last_role = len(role_values) + 1
    roles_list = DataValidation(type="list", formula1=f"=Roles!$A$2:$A${last_role}",
                                allow_blank=False, showErrorMessage=True,
                                errorTitle="Not a role",
                                error="Pick a role from the list (see the Roles sheet).")
    sheet.add_data_validation(roles_list)
    roles_list.add(f"H2:H{sheet.max_row}")
    reviewed = DataValidation(type="list", formula1='"Y"', allow_blank=True)
    pages_sheet.add_data_validation(reviewed)
    reviewed.add(f"E2:E{pages_sheet.max_row}")

    widths = {"A": 34, "B": 6, "C": 6, "D": 7, "E": 28, "F": 40, "G": 16, "H": 16}
    for letter, width in widths.items():
        sheet.column_dimensions[letter].width = width
    for letter in "IJKL":
        sheet.column_dimensions[letter].hidden = True
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:H{sheet.max_row}"
    for row in sheet.iter_rows(min_row=2, max_col=6):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (5, 6))
    pages_sheet.column_dimensions["A"].width = 44
    pages_sheet.freeze_panes = "A2"

    LABEL_DIR.mkdir(parents=True, exist_ok=True)
    target = WORKBOOK
    if target.exists():
        # Never overwrite a sheet someone may have spent an afternoon on.
        target = LABEL_DIR / "column_roles.new.xlsx"
        print(f"  {WORKBOOK.name} already exists and was left untouched")
    book.save(target)
    print(f"wrote {target.relative_to(ROOT)}: {len(chosen)} page(s), {rows_written} column(s)")
    return target


# ---------------------------------------------------------------------------
# score
# ---------------------------------------------------------------------------


def load_labels() -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    if not WORKBOOK.is_file():
        return []
    book = load_workbook(WORKBOOK, data_only=True)
    reviewed = {(str(r[0]), int(r[1])) for r in book["Pages"].iter_rows(min_row=2, values_only=True)
                if r[0] and str(r[4] or "").strip().upper() == "Y"}
    labels = []
    for row in book["Labels"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        key = (str(row[0]), int(row[1]))
        if key not in reviewed:
            continue
        labels.append({"file": key[0], "page": key[1], "heading": row[4] or "",
                       "label": str(row[7] or "").strip(),
                       "x": (float(row[8]), float(row[9])),
                       "y": (float(row[10]), float(row[11]))})
    return labels


def score(json_out: Path | None = None) -> dict[str, Any]:
    """Score today's code against the reviewed labels.

    Two numbers, deliberately kept apart, because they mean different things:

    * **wrong roles** - the system named a column, and named it wrongly. This is
      the dangerous class: a consumer reading `opening_qty` gets the receipt
      column, and nothing downstream says so. It should only ever fall.
    * **missed roles** - the column has a role and the system left it unknown.
      That costs review effort but misleads no one.
    """
    from pipeline import process_file

    labels = load_labels()
    summary: dict[str, Any] = {"labelled_pages": 0, "labelled_columns": len(labels)}
    if not labels:
        summary["note"] = ("no reviewed pages: mark pages Y in the Pages sheet of "
                           f"{WORKBOOK.relative_to(ROOT)}")
        if json_out:
            json_out.write_text(json.dumps(summary, indent=1), encoding="utf-8")
        return summary

    by_file: dict[str, list[dict]] = defaultdict(list)
    for label in labels:
        by_file[label["file"]].append(label)

    canonical = correct = missed = wrong = not_found = 0
    assigned = 0
    per_page: dict[str, dict[str, int]] = {}
    mistakes: list[str] = []
    for name, file_labels in sorted(by_file.items()):
        result = process_file(CORPUS / name)
        pages = {label["page"] for label in file_labels}
        predicted = {p: predicted_columns(result.pages[p - 1].document)
                     if result.pages[p - 1].document else [] for p in pages}
        for label in file_labels:
            key = f"{name} p{label['page']}"
            stats = per_page.setdefault(key, {"correct": 0, "missed": 0, "wrong": 0,
                                              "not_found": 0})
            hit = match(label["x"], label["y"], predicted[label["page"]])
            truth = label["label"]
            has_role = truth not in ("unknown", MERGED, NOT_A_COLUMN)
            if has_role:
                canonical += 1
            if hit is None:
                if has_role:
                    not_found += 1
                    stats["not_found"] += 1
                    mistakes.append(f"{key}: {truth} column ({label['heading']!r}) no longer found")
                continue
            if hit.role != "unknown":
                assigned += 1
            if has_role and hit.role == truth:
                correct += 1
                stats["correct"] += 1
            elif hit.role == "unknown":
                if has_role:
                    missed += 1
                    stats["missed"] += 1
            else:
                wrong += 1
                stats["wrong"] += 1
                mistakes.append(f"{key}: {label['heading']!r} is {truth}, system says {hit.role}")

    summary.update({
        "labelled_pages": len(per_page),
        "columns_with_a_role": canonical,
        "correct": correct,
        "missed": missed,
        "not_found": not_found,
        "roles_assigned": assigned,
        "wrong": wrong,
        "per_page": per_page,
        "mistakes": mistakes,
    })
    if json_out:
        json_out.write_text(json.dumps(summary, indent=1), encoding="utf-8")
    return summary


def print_score(summary: dict[str, Any]) -> None:
    if not summary.get("labelled_columns"):
        print(summary.get("note", "no labels"))
        return
    c, n = summary["correct"], summary["columns_with_a_role"]
    w, a = summary["wrong"], summary["roles_assigned"]
    print(f"labelled pages : {summary['labelled_pages']}")
    print(f"correct roles  : {c} of {n} columns that have a role"
          + (f" ({c / n:.1%})" if n else ""))
    print(f"missed         : {summary['missed']} left unknown, "
          f"{summary['not_found']} no longer found")
    print(f"WRONG roles    : {w} of {a} roles the system assigned"
          + (f" ({w / a:.1%})" if a else ""))
    for line in summary["mistakes"][:40]:
        print(f"   {line}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", required=True)
    make_cmd = sub.add_parser("make", help="write the pre-filled labelling workbook")
    make_cmd.add_argument("--pages", type=int, default=40)
    score_cmd = sub.add_parser("score", help="score the current code against reviewed labels")
    score_cmd.add_argument("--json", type=Path)
    args = parser.parse_args(argv)

    if args.command == "make":
        make(args.pages)
    else:
        print_score(score(args.json))
    return 0


if __name__ == "__main__":
    sys.exit(main())
